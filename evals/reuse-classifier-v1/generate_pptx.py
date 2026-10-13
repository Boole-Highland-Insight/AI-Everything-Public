#!/usr/bin/env python3
"""Generate the PPTX portion of the 100-instance reuse benchmark."""

from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "instances" / "pptx"
PARTIAL_MANIFEST = ROOT / ".build" / "manifest.non-pptx.jsonl"
FINAL_MANIFEST = ROOT / "labels.jsonl"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def color(value: str) -> RGBColor:
    value = value.removeprefix("#")
    return RGBColor.from_string(value)


def add_text(slide, text: str, left: float, top: float, width: float, height: float,
             size: int, bold: bool = False, text_color: str = "#263238") -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "PingFang SC"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color(text_color)


def make_presentation(spec: dict[str, object], path: Path) -> None:
    presentation = Presentation()
    presentation.core_properties.identifier = str(spec["id"])
    presentation.core_properties.title = str(spec["title"])
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color({
        "completion": "#FAF8F3",
        "optimization": "#F5F8F6",
        "none": "#F6F8FA",
        "duplicate": "#F6F8FA",
    }[str(spec["label"])])
    add_text(slide, str(spec["title"]), 0.88, 0.55, 11.55, 0.75, 28, True, "#23313A")
    add_text(slide, "\n\n".join(spec["lines"]), 0.96, 1.75, 10.95, 4.05, 17, False, "#425466")
    add_text(slide, "咔哒分类评测夹具", 0.96, 6.62, 3.75, 0.3, 10, False, "#82909C")
    presentation.save(path)


def specs() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for index in range(1, 6):
        explicit = index <= 2
        rows.append({
            "id": f"ppt-completion-{index:03d}",
            "filename": f"季度计划_未完成_{index}.pptx" if explicit else f"季度计划_{index}.pptx",
            "label": "completion", "cue": "explicit" if explicit else "semantic",
            "title": "季度计划 未完成" if explicit else f"季度计划 {index}",
            "lines": ["目标与现状已填写", "资源安排：产品、研发、设计", "风险与缓解措施：待补充", "下一页应包含验收标准，但尚未创建"],
            "rationale": "演示文稿缺少风险细节和验收部分。",
        })
    for index in range(1, 5):
        explicit = index <= 2
        rows.append({
            "id": f"ppt-optimization-{index:03d}",
            "filename": f"产品汇报草稿_{index}.pptx" if explicit else f"产品汇报_{index}.pptx",
            "label": "optimization", "cue": "explicit" if explicit else "semantic",
            "title": "产品汇报草稿" if explicit else f"产品汇报 {index}",
            "lines": ["目标、结果、风险和下一步均已覆盖", "当前每段文字较长，重点层级不够清楚", "事实完整，适合重新组织表达"],
            "rationale": "内容完整，但信息层级和措辞需要打磨。",
        })
    for index in range(1, 10):
        rows.append({
            "id": f"ppt-none-{index:03d}", "filename": f"完整项目汇报_{index:02d}.pptx",
            "label": "none", "cue": "semantic", "title": f"项目汇报 {index}",
            "lines": ["目标：提升检索准确率并保留来源", "结果：核心流程通过验证", "风险：模型不可用时启用规则回退", "结论：本阶段目标已经完成"],
            "rationale": "目标、结果、风险和结论完整。",
        })
    rows.append({
        "id": "ppt-duplicate-001", "filename": "交付汇报_原件.pptx",
        "label": "duplicate", "cue": "exact_hash", "title": "交付汇报",
        "lines": ["范围：搜索、来源展示和恢复", "结果：验收项全部通过", "结论：同意交付"],
        "rationale": "与同组 PPTX 字节完全相同。", "duplicate_group": "dup-pptx-01",
    })
    return rows


def main() -> None:
    if not PARTIAL_MANIFEST.is_file():
        raise SystemExit("先运行 generate_non_pptx.py")
    shutil.rmtree(OUTPUT, ignore_errors=True)
    OUTPUT.mkdir(parents=True)
    ppt_rows: list[dict[str, object]] = []
    for spec in specs():
        output = OUTPUT / str(spec["filename"])
        make_presentation(spec, output)
        ppt_rows.append({
            "id": spec["id"], "relative_path": f"instances/pptx/{spec['filename']}",
            "media_type": "pptx", "extension": "pptx", "label": spec["label"],
            "cue": spec["cue"], "rationale": spec["rationale"],
            "duplicate_group": spec.get("duplicate_group"), "sha256": sha256(output),
        })
    source = OUTPUT / "交付汇报_原件.pptx"
    duplicate = OUTPUT / "交付汇报_副本.pptx"
    shutil.copyfile(source, duplicate)
    ppt_rows.append({**ppt_rows[-1], "id": "ppt-duplicate-002",
                     "relative_path": "instances/pptx/交付汇报_副本.pptx",
                     "sha256": sha256(duplicate)})
    base_rows = [json.loads(line) for line in PARTIAL_MANIFEST.read_text(encoding="utf-8").splitlines() if line]
    all_rows = base_rows + ppt_rows
    FINAL_MANIFEST.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in all_rows), encoding="utf-8")
    print(f"generated {len(ppt_rows)} PPTX instances and {len(all_rows)} total labels")


if __name__ == "__main__":
    main()
